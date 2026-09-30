import os
from io import BytesIO
from datetime import date

import streamlit as st
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    .main-title {
        text-align: center;
        font-size: 42px;
        font-weight: 800;
        margin-bottom: 5px;
    }

    .subtitle {
        text-align: center;
        color: #9CA3AF;
        font-size: 17px;
        margin-bottom: 30px;
    }

    .notice {
        padding: 15px;
        border-radius: 10px;
        background: #172554;
        border: 1px solid #2563EB;
        margin-top: 20px;
        margin-bottom: 20px;
    }

    .document-box {
        padding: 20px;
        border-radius: 10px;
        background: #111827;
        border: 1px solid #374151;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# TITLE
# ============================================================

st.markdown(
    '<div class="main-title">⚖️ LegalEase</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">'
    'AI-Assisted Legal Document Drafting System'
    '</div>',
    unsafe_allow_html=True
)


# ============================================================
# GEMINI
# ============================================================

def get_gemini_client():

    api_key = None

    # Streamlit Cloud secrets
    try:
        api_key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        pass

    # Environment variable
    if not api_key:
        api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:
        return None

    try:
        from google import genai

        return genai.Client(api_key=api_key)

    except Exception:
        return None


def generate_with_gemini(
    document_type,
    parties,
    terms,
    effective_date
):

    client = get_gemini_client()

    if client is None:
        return None

    prompt = f"""
You are LegalEase, an AI-assisted legal document drafting system.

Create a professional legal document draft using ONLY the
information provided below.

IMPORTANT RULES:

- Do not invent names.
- Do not invent addresses.
- Do not invent dates.
- Do not invent monetary amounts.
- Do not invent jurisdictions.
- Missing information must be written as [INSERT INFORMATION].
- Do not claim that the document is legally valid.
- Do not claim that it is legally enforceable.
- Use formal professional legal language.
- Use numbered sections.
- Do not use Markdown code blocks.
- End with a drafting notice.

DOCUMENT TYPE:
{document_type}

PARTIES:
{parties}

TERMS:
{terms}

EFFECTIVE DATE:
{effective_date}

Use relevant sections such as:

1. Purpose
2. Parties
3. Definitions
4. Scope
5. Terms and Conditions
6. Payment
7. Confidentiality
8. Intellectual Property
9. Term and Termination
10. Liability
11. Governing Law
12. Dispute Resolution
13. Severability
14. Entire Agreement
15. Signatures

Only include relevant sections.

At the end write:

DRAFTING NOTICE

This document is an AI-assisted draft and should be reviewed
by a qualified legal professional before signing, filing,
or relying upon it.
"""

    try:

        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt
        )

        return response.text

    except Exception as error:

        st.warning(
            f"Gemini generation failed: {error}"
        )

        return None


# ============================================================
# DEMO FALLBACK
# ============================================================

def demo_document(
    document_type,
    parties,
    terms,
    effective_date
):

    return f"""
{document_type.upper()}

EFFECTIVE DATE

{effective_date}

1. PARTIES

This document is entered into between the following parties:

{parties}

2. PURPOSE

The purpose of this document is to establish the terms and
conditions agreed upon by the parties.

3. TERMS AND CONDITIONS

{terms}

4. CONFIDENTIALITY

The parties shall maintain confidentiality regarding
confidential information exchanged in connection with this
agreement, subject to applicable law.

5. TERM AND TERMINATION

The term and termination conditions shall be determined by
the parties based on their agreed requirements.

6. LIABILITY

Each party shall be responsible for its obligations under
this document, subject to applicable law.

7. GOVERNING LAW

[INSERT GOVERNING LAW / JURISDICTION]

8. DISPUTE RESOLUTION

[INSERT DISPUTE RESOLUTION MECHANISM]

9. ENTIRE AGREEMENT

This document represents the understanding of the parties
regarding the subject matter described herein.

10. SIGNATURES


PARTY 1

Name: ______________________________

Signature: _________________________

Date: ______________________________


PARTY 2

Name: ______________________________

Signature: _________________________

Date: ______________________________


DRAFTING NOTICE

This document is an AI-assisted demo draft and should be
reviewed by a qualified legal professional before signing,
filing, or relying upon it.
"""


# ============================================================
# DOCUMENT EXPORT
# ============================================================

def create_txt(text):

    return text.encode("utf-8")


def create_docx(text):

    document = Document()

    for line in text.splitlines():

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)

        run.font.name = "Times New Roman"

    output = BytesIO()

    document.save(output)

    return output.getvalue()


def create_pdf(text):

    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=50,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    title_style.alignment = TA_CENTER

    body_style = styles["BodyText"]

    story = []

    lines = text.splitlines()

    first_line = True

    for line in lines:

        line = line.strip()

        if not line:

            story.append(Spacer(1, 8))

            continue

        safe_line = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if first_line:

            story.append(
                Paragraph(
                    safe_line,
                    title_style
                )
            )

            first_line = False

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style
                )
            )

    pdf.build(story)

    return output.getvalue()


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header("⚙️ LegalEase")

    st.write(
        "AI-assisted legal document drafting."
    )

    st.divider()

    gemini_available = get_gemini_client() is not None

    if gemini_available:

        st.success("Gemini API configured")

    else:

        st.info(
            "Demo mode\n\n"
            "Add GEMINI_API_KEY to enable AI generation."
        )

    st.divider()

    st.caption(
        "LegalEase is an AI-assisted drafting tool. "
        "Generated documents require professional legal review."
    )


# ============================================================
# INPUT SECTION
# ============================================================

st.subheader("📝 Create Legal Document")

col1, col2 = st.columns(2)

with col1:

    document_type = st.selectbox(
        "Document Type",
        [
            "Freelance Work Contract",
            "Non-Disclosure Agreement",
            "Employment Agreement",
            "Service Agreement",
            "Rental Agreement",
            "Partnership Agreement",
            "Business Agreement",
            "Custom Legal Document"
        ]
    )

    parties = st.text_area(
        "Parties",
        placeholder=(
            "Example:\n"
            "John Doe, Service Provider\n"
            "ABC Technologies Pvt Ltd, Client"
        ),
        height=130
    )

with col2:

    effective_date = st.date_input(
        "Effective Date",
        value=date.today()
    )

    terms = st.text_area(
        "Terms and Conditions",
        placeholder=(
            "Example:\n"
            "Payment within 30 days; "
            "Confidentiality; "
            "Project duration 6 months"
        ),
        height=130
    )


# ============================================================
# GENERATE
# ============================================================

st.divider()

generate_button = st.button(
    "⚡ Generate Legal Document",
    type="primary",
    use_container_width=True
)


if generate_button:

    if not parties.strip():

        st.error("Please enter the parties.")

        st.stop()

    if not terms.strip():

        st.error("Please enter the terms and conditions.")

        st.stop()

    with st.spinner("Generating legal document..."):

        generated = generate_with_gemini(
            document_type,
            parties,
            terms,
            effective_date
        )

        if generated:

            st.session_state["document"] = generated

        else:

            st.session_state["document"] = demo_document(
                document_type,
                parties,
                terms,
                effective_date
            )


# ============================================================
# DOCUMENT EDITOR
# ============================================================

if "document" in st.session_state:

    st.divider()

    st.subheader("📄 Generated Document")

    edited_document = st.text_area(
        "Edit your document",
        value=st.session_state["document"],
        height=600
    )

    st.session_state["document"] = edited_document

    st.divider()

    st.subheader("👁️ Preview")

    st.markdown(
        '<div class="document-box">',
        unsafe_allow_html=True
    )

    for line in edited_document.splitlines():

        line = line.strip()

        if not line:
            st.write("")
            continue

        if line[0:2] == "1." or line[0:2] == "2.":

            st.markdown(f"**{line}**")

        else:

            st.write(line)

    st.markdown(
        "</div>",
        unsafe_allow_html=True
    )

    # ========================================================
    # DOWNLOAD
    # ========================================================

    st.divider()

    st.subheader("⬇️ Export Document")

    col1, col2, col3 = st.columns(3)

    with col1:

        st.download_button(
            label="📄 Download TXT",
            data=create_txt(edited_document),
            file_name="legalease_document.txt",
            mime="text/plain",
            use_container_width=True
        )

    with col2:

        st.download_button(
            label="📝 Download DOCX",
            data=create_docx(edited_document),
            file_name="legalease_document.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True
        )

    with col3:

        st.download_button(
            label="📕 Download PDF",
            data=create_pdf(edited_document),
            file_name="legalease_document.pdf",
            mime="application/pdf",
            use_container_width=True
        )


# ============================================================
# LEGAL NOTICE
# ============================================================

st.divider()

st.warning(
    "⚠️ Legal Notice: LegalEase generates AI-assisted draft "
    "documents. It does not provide legal advice. "
    "Always have a qualified legal professional review "
    "the document before signing or relying upon it."
)
