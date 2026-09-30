import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)

from utils.sanitizer import sanitize_text, escape_html


def parse_lines(text: str):
    text = sanitize_text(text)
    return text.split("\n")


def format_html_preview(text: str) -> str:
    lines = parse_lines(text)

    html_parts = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            html_parts.append("<div style='height:10px'></div>")
            continue

        escaped = escape_html(stripped)

        if stripped.startswith("# "):
            html_parts.append(
                f"<h1>{escape_html(stripped[2:])}</h1>"
            )

        elif stripped.startswith("## "):
            html_parts.append(
                f"<h2>{escape_html(stripped[3:])}</h2>"
            )

        elif stripped.startswith("### "):
            html_parts.append(
                f"<h3>{escape_html(stripped[4:])}</h3>"
            )

        elif re.match(r"^\d+\.\s+", stripped):
            html_parts.append(
                f"<p><strong>{escaped}</strong></p>"
            )

        elif stripped.startswith("- "):
            html_parts.append(
                f"<p style='margin-left:25px'>• "
                f"{escape_html(stripped[2:])}</p>"
            )

        else:
            html_parts.append(
                f"<p>{escaped}</p>"
            )

    return "\n".join(html_parts)


def _set_cell_text(cell, text, bold=False):
    cell.text = ""

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _apply_docx_font(document):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)


def format_docx(
    text: str,
    document_type: str,
    parties: str = "",
    dates: str = "",
    terms: str = "",
    logo_path=None
):
    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Logo
    if logo_path:
        logo = Path(logo_path)

        if logo.exists():
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run()
            run.add_picture(
                str(logo),
                width=Inches(1.2)
            )

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        sanitize_text(document_type).upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    # Metadata table
    metadata = [
        ("Document Type", document_type),
        ("Effective Date", dates),
        ("Parties", parties),
    ]

    table = document.add_table(
        rows=len(metadata),
        cols=2
    )

    table.style = "Table Grid"

    for row, (label, value) in zip(
        table.rows,
        metadata
    ):
        _set_cell_text(
            row.cells[0],
            label,
            bold=True
        )

        _set_cell_text(
            row.cells[1],
            value
        )

    document.add_paragraph()

    # Terms table
    if terms.strip():
        heading = document.add_paragraph()

        run = heading.add_run(
            "Key Terms and Conditions"
        )

        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)

        terms_list = [
            item.strip()
            for item in re.split(
                r"[;\n]+",
                terms
            )
            if item.strip()
        ]

        if terms_list:
            terms_table = document.add_table(
                rows=1,
                cols=2
            )

            terms_table.style = "Table Grid"

            _set_cell_text(
                terms_table.rows[0].cells[0],
                "No.",
                bold=True
            )

            _set_cell_text(
                terms_table.rows[0].cells[1],
                "Term",
                bold=True
            )

            for index, term in enumerate(
                terms_list,
                start=1
            ):
                row = terms_table.add_row()

                _set_cell_text(
                    row.cells[0],
                    index
                )

                _set_cell_text(
                    row.cells[1],
                    term
                )

            document.add_paragraph()

    # Main document
    for line in parse_lines(text):
        stripped = line.strip()

        if not stripped:
            document.add_paragraph()
            continue

        if stripped.startswith("# "):
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                stripped[2:]
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        elif stripped.startswith("## "):
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                stripped[3:]
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)

        elif stripped.startswith("### "):
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                stripped[4:]
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            run = paragraph.add_run(
                stripped[2:]
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        elif re.match(
            r"^\d+\.\s+",
            stripped
        ):
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                stripped
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        else:
            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                stripped
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Footer
    footer = section.footer

    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "LegalEase - AI-Assisted Legal Document Draft"
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(8)

    _apply_docx_font(document)

    output = BytesIO()

    document.save(output)

    return output.getvalue()


def _pdf_header_footer(
    canvas,
    doc,
    document_type
):
    canvas.saveState()

    width, height = A4

    canvas.setFont(
        "Helvetica",
        8
    )

    canvas.drawString(
        20 * mm,
        height - 12 * mm,
        "LegalEase"
    )

    canvas.drawRightString(
        width - 20 * mm,
        height - 12 * mm,
        sanitize_text(document_type)
    )

    canvas.drawCentredString(
        width / 2,
        10 * mm,
        f"LegalEase - Page {doc.page}"
    )

    canvas.restoreState()


def format_pdf(
    text: str,
    document_type: str,
    parties: str = "",
    dates: str = "",
    terms: str = "",
    logo_path=None
):
    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=18 * mm
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        alignment=TA_CENTER,
        spaceAfter=12
    )

    heading_style = ParagraphStyle(
        "LegalHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        spaceBefore=10,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        alignment=TA_LEFT,
        spaceAfter=6
    )

    small_style = ParagraphStyle(
        "LegalSmall",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=12
    )

    story = []

    # Logo
    if logo_path:
        logo = Path(logo_path)

        if logo.exists():
            from reportlab.platypus import Image

            image = Image(
                str(logo),
                width=30 * mm,
                height=30 * mm
            )

            image.hAlign = "CENTER"

            story.append(image)
            story.append(Spacer(1, 5 * mm))

    story.append(
        Paragraph(
            escape_html(
                document_type.upper()
            ),
            title_style
        )
    )

    # Metadata
    metadata = [
        [
            Paragraph("<b>Document Type</b>", small_style),
            Paragraph(
                escape_html(document_type),
                small_style
            )
        ],
        [
            Paragraph("<b>Effective Date</b>", small_style),
            Paragraph(
                escape_html(dates),
                small_style
            )
        ],
        [
            Paragraph("<b>Parties</b>", small_style),
            Paragraph(
                escape_html(parties),
                small_style
            )
        ]
    ]

    metadata_table = Table(
        metadata,
        colWidths=[40 * mm, 120 * mm]
    )

    metadata_table.setStyle(
        TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
    )

    story.append(metadata_table)
    story.append(Spacer(1, 8 * mm))

    # Terms
    if terms.strip():
        story.append(
            Paragraph(
                "Key Terms and Conditions",
                heading_style
            )
        )

        terms_list = [
            item.strip()
            for item in re.split(
                r"[;\n]+",
                terms
            )
            if item.strip()
        ]

        term_data = [
            [
                Paragraph("<b>No.</b>", small_style),
                Paragraph("<b>Term</b>", small_style)
            ]
        ]

        for index, term in enumerate(
            terms_list,
            start=1
        ):
            term_data.append(
                [
                    Paragraph(
                        str(index),
                        small_style
                    ),
                    Paragraph(
                        escape_html(term),
                        small_style
                    )
                ]
            )

        terms_table = Table(
            term_data,
            colWidths=[15 * mm, 145 * mm]
        )

        terms_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )

        story.append(terms_table)
        story.append(Spacer(1, 8 * mm))

    # Main content
    for line in parse_lines(text):
        stripped = line.strip()

        if not stripped:
            story.append(Spacer(1, 3 * mm))
            continue

        if stripped.startswith("# "):
            story.append(
                Paragraph(
                    escape_html(stripped[2:]),
                    title_style
                )
            )

        elif stripped.startswith("## "):
            story.append(
                Paragraph(
                    escape_html(stripped[3:]),
                    heading_style
                )
            )

        elif stripped.startswith("### "):
            story.append(
                Paragraph(
                    f"<b>{escape_html(stripped[4:])}</b>",
                    heading_style
                )
            )

        elif stripped.startswith("- "):
            story.append(
                Paragraph(
                    f"• {escape_html(stripped[2:])}",
                    body_style
                )
            )

        else:
            story.append(
                Paragraph(
                    escape_html(stripped),
                    body_style
                )
            )

    document.build(
        story,
        onFirstPage=lambda canvas, doc:
            _pdf_header_footer(
                canvas,
                doc,
                document_type
            ),
        onLaterPages=lambda canvas, doc:
            _pdf_header_footer(
                canvas,
                doc,
                document_type
            )
    )

    return output.getvalue()